"""Unit tests for create_spreadsheet, author_code, and python_verify arms."""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pipeline import author_code, create_spreadsheet, python_verify


class CreateSpreadsheetTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.out = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_write_xlsx(self) -> None:
        with patch.dict(os.environ, {"EMPIRE_EVE_OUTPUT_DIR": str(self.out)}):
            res = create_spreadsheet.create_spreadsheet(
                filename="report",
                headers=["name", "age"],
                rows=[["ada", 36], ["grace", 44]],
            )
        self.assertTrue(res["ok"], res.get("error"))
        self.assertTrue(Path(res["path"]).is_file())
        self.assertTrue(res["path"].endswith(".xlsx"))

    def test_formula_injection_neutralized(self) -> None:
        with patch.dict(os.environ, {"EMPIRE_EVE_OUTPUT_DIR": str(self.out)}):
            res = create_spreadsheet.create_spreadsheet(
                filename="inject",
                headers=["a"],
                rows=[["=CMD()"]],
            )
        self.assertTrue(res["ok"], res.get("error"))
        # Re-open and confirm the cell is stored as inert text.
        from openpyxl import load_workbook

        wb = load_workbook(res["path"])
        ws = wb.active
        cell = ws.cell(row=2, column=1).value
        self.assertTrue(str(cell).startswith("'"))

    def test_missing_headers(self) -> None:
        res = create_spreadsheet.create_spreadsheet(filename="x", headers=[], rows=[])
        self.assertFalse(res["ok"])

    def test_external_link_neutralized(self) -> None:
        with patch.dict(os.environ, {"EMPIRE_EVE_OUTPUT_DIR": str(self.out)}):
            res = create_spreadsheet.create_spreadsheet(
                filename="link",
                headers=["a"],
                rows=[["=HYPERLINK(\"http://evil\",\"x\")"]],
            )
        self.assertTrue(res["ok"], res.get("error"))


class AuthorCodeTests(unittest.TestCase):
    def test_path_escape_rejected(self) -> None:
        res = author_code.apply_patch(
            worktree=str(Path("C:/tmp/wt")),
            relative_path="../secret.py",
            content="x",
        )
        self.assertFalse(res["ok"])

    def test_credential_path_rejected(self) -> None:
        res = author_code.apply_patch(
            worktree=str(Path("C:/tmp/wt")),
            relative_path=".env",
            content="KEY=1",
        )
        self.assertFalse(res["ok"])
        self.assertIn("credential", res["error"])

    def test_abs_path_rejected(self) -> None:
        self.assertTrue(author_code.re_abs("C:\\x\\y.py"))
        self.assertTrue(author_code.re_abs("/etc/passwd"))
        self.assertFalse(author_code.re_abs("src/main.py"))


class PythonVerifyTests(unittest.TestCase):
    def test_syntax_ok_on_valid_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "good.py").write_text("def f():\n    return 1\n", encoding="utf-8")
            res = python_verify.verify(worktree=str(root), run_tests=False)
            self.assertTrue(res["ok"], res)
            self.assertTrue(res["report"]["syntax"]["ok"])

    def test_syntax_fails_on_bad_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "bad.py").write_text("def f(:\n", encoding="utf-8")
            res = python_verify.verify(worktree=str(root), run_tests=False)
            self.assertFalse(res["ok"])
            self.assertFalse(res["report"]["syntax"]["ok"])

    def test_missing_worktree(self) -> None:
        res = python_verify.verify(worktree=str(Path("C:/does/not/exist")), run_tests=False)
        self.assertFalse(res["ok"])


if __name__ == "__main__":
    unittest.main()
